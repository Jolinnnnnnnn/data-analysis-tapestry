"""01 - Inventory the raw workbook: sheets, visibility, dimensions, merged cells,
header offsets, comments, and full dump of Column Definitions.

Output: outputs/tables/01_workbook_inventory.md
"""
from pathlib import Path
import openpyxl
import pandas as pd

ROOT = Path(__file__).resolve().parents[1]
RAW = ROOT / "data/raw/TYLY_data_sample.xlsx"
OUT = ROOT / "outputs/tables/01_workbook_inventory.md"

wb = openpyxl.load_workbook(RAW, data_only=False)
lines = ["# Workbook inventory", f"Source: `{RAW.relative_to(ROOT)}`", ""]
lines.append("| sheet | state | max_row | max_col | merged ranges | comments | formulas |")
lines.append("|---|---|---|---|---|---|---|")
for ws in wb.worksheets:
    n_comments = sum(1 for row in ws.iter_rows() for c in row if c.comment)
    n_formula = sum(1 for row in ws.iter_rows() for c in row
                    if isinstance(c.value, str) and c.value.startswith("="))
    lines.append(f"| {ws.title} | {ws.sheet_state} | {ws.max_row} | {ws.max_column} | "
                 f"{list(map(str, ws.merged_cells.ranges)) or '-'} | {n_comments} | {n_formula} |")
lines.append("")
lines.append(f"Defined names: {list(wb.defined_names.keys()) if hasattr(wb.defined_names, 'keys') else wb.defined_names}")
lines.append("")

for ws in wb.worksheets:
    lines.append(f"## Sheet `{ws.title}` — first 12 raw rows (header-offset check)")
    lines.append("```")
    for i, row in enumerate(ws.iter_rows(min_row=1, max_row=12, values_only=True), 1):
        lines.append(f"{i:>3}: {row}")
    lines.append("```")
    # autofilter / freeze panes / hidden rows/cols
    hidden_rows = [r for r, d in ws.row_dimensions.items() if d.hidden]
    hidden_cols = [c for c, d in ws.column_dimensions.items() if d.hidden]
    lines.append(f"- freeze_panes: {ws.freeze_panes}; autofilter: {ws.auto_filter.ref}; "
                 f"hidden rows: {hidden_rows[:20]} (n={len(hidden_rows)}); hidden cols: {hidden_cols}")
    lines.append("")

# Full dump of Column Definitions
defs = pd.read_excel(RAW, sheet_name="Column Definitions", header=None)
lines.append("## Column Definitions (full dump)")
lines.append("```")
for i, r in defs.iterrows():
    lines.append(" | ".join("" if pd.isna(v) else str(v) for v in r.values))
lines.append("```")

OUT.write_text("\n".join(lines))
print("\n".join(lines))
